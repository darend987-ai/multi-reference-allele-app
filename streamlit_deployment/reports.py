from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill
from core import compare, excluded

def export(jobs, samples):
    book = openpyxl.Workbook()
    book.remove(book.active)
    def sheet(name, records):
        ws = book.create_sheet(name)
        if not records:
            ws.append(['No records'])
            return
        headers = list(records[0])
        ws.append(headers)
        for record in records:
            # Explicit text cells prevent sample names/labels becoming formulas.
            for col, key in enumerate(headers, 1):
                value = record.get(key)
                cell = ws.cell(ws.max_row + 1, col) if col == 1 else ws.cell(ws.max_row, col)
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = 's'
                if key in ('recovery', 'unique_recovery') or key.endswith((' recovery', ' unique_recovery')):
                    cell.number_format = '0.0%'
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='24476A')
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(55, max(14, max(len(str(c.value or '')) for c in col)+2))
    master, ref_summary, assignments = [], [], []
    for ev, refs, count in jobs:
        if len(refs) > count:
            raise ValueError(f'{ev.name}: {len(refs)} references exceed the {count} contributor slots. Increase the contributor count or select fewer references.')
    max_references = max((count for _, _, count in jobs), default=0)
    for index, (ev, refs, count) in enumerate(jobs, 1):
        detail, attribution, totals, foreign, shared = compare(ev, refs)
        summary = {'Mixture': ev.name, 'Mixture ID': ev.id, 'Contributors': count,
                   'References Evaluated': len(refs)}
        for slot in range(max_references):
            summary[f'Contributor {slot + 1}'] = refs[slot].name if slot < len(refs) else None
        for slot in range(max_references):
            t = totals[slot] if slot < len(totals) else None
            for category in ('observed', 'expected', 'missing', 'unique_observed',
                             'unique_expected', 'recovery', 'unique_recovery'):
                summary[f'Contributor {slot + 1} {category}'] = t[category] if t else None
        summary.update({'Foreign Alleles': foreign, 'Shared Alleles': shared, 'Detail Sheet': f'PL_{index}'})
        master.append(summary)
        for i, r in enumerate(refs, 1):
            assignments.append({'Mixture ID': ev.id, 'Reference Label': f'R{i}', 'Reference': r.name, 'Reference ID': r.id})
        ref_summary.extend({'Mixture': ev.name, 'Mixture ID': ev.id, **t} for t in totals)
        sheet(f'PL_{index}', detail)
        sheet(f'ALLELES_{index}', attribution)
    sheet('Master_Summary', master)
    book['Master_Summary'].freeze_panes = 'C2'
    sheet('Reference_Summary', ref_summary)
    sheet('Assignments', assignments)
    sheet('Import_Audit', [{'Sample ID': s.id, 'Sample': s.name, 'Source': s.source, 'Source Row': s.row,
                          'Marker': locus, 'Raw Labels': ', '.join(values),
                          'Excluded': excluded(locus) if not locus.startswith('_') else False}
                         for s in samples for locus, values in s.raw.items()])
    sheet('Import_Warnings', [{'Sample ID': s.id, 'Warning': w} for s in samples for w in s.warnings])
    book._sheets.insert(0, book._sheets.pop(book.sheetnames.index('Reference_Summary')))
    book._sheets.insert(0, book._sheets.pop(book.sheetnames.index('Master_Summary')))
    out = BytesIO()
    book.save(out)
    return out.getvalue()
